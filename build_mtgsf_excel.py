import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

headers = [
    'Sıra No', 'T.C. Kimlik No', 'Adı Soyadı', 'Öğrencinin Üniversitesi / Fakültesi / Programı',
    'Öğrenci YKS Yılı', 'Öğrenci Puan Türü', 'Öğrenci Puanı',
    'OKÜ Taban Puan Yılı', 'OKÜ Taban Puan Türü', 'OKÜ Taban Puanı',
    'Yatay Geçiş Yapılan Bölüm / Program', 'Sonuç (Asıl / Yedek)', 'Kayıt Olacağı Sınıf'
]

rows = [
    ['1.', '11063424696', 'Muhammed Emin BADEM', 'Siirt Üniversitesi / Eğitim Fakültesi / Türkçe Öğretmenliği', '2025', 'EA', '319,01982', '2025', 'EA', '318,46604', 'İç Mimarlık ve Çevre Tasarımı Bölümü', 'ASIL', '1'],
    ['2.', '59476341346', 'Zühre KARAKUM', 'Başkent Üniversitesi / Güzel Sanatlar Tasarım ve Mimarlık Fakültesi / İç Mimarlık ve Çevre Tasarımı', '2022', 'EA', '384,77019', '2022', 'EA', '327,20219', 'İç Mimarlık ve Çevre Tasarımı Bölümü', 'ASIL', '1'],
    ['3.', '11300839498', 'Aylin KORKUT', 'Girne Amerikan Üniversitesi / Beşeri Bilimler Fakültesi / Psikoloji', '2025', 'EA', '320,85667', '2025', 'EA', '318,46604', 'İç Mimarlık ve Çevre Tasarımı Bölümü', 'ASIL', '1']
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'İç Mimarlık ve Çevre Tasarımı'

faculty_banner_fill = PatternFill(start_color='0D233A', end_color='0D233A', fill_type='solid')
faculty_banner_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')

dept_banner_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
dept_banner_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

col_header_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
col_header_font = Font(name='Calibri', size=10, bold=True, color='000000')

data_font = Font(name='Calibri', size=10, color='000000')
zebra_fill = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')

thin_border = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

thick_bottom = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='medium', color='1F497D')
)

# Faculty Title
ws.merge_cells('A1:M1')
top_cell = ws['A1']
top_cell.value = 'OSMANİYE KORKUT ATA ÜNİVERSİTESİ - MİMARLIK TASARIM VE GÜZEL SANATLAR FAKÜLTESİ'
top_cell.fill = faculty_banner_fill
top_cell.font = faculty_banner_font
top_cell.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 28

ws.merge_cells('A2:M2')
sub_top = ws['A2']
sub_top.value = '2026-2027 EĞİTİM-ÖĞRETİM YILI GÜZ YARIYILI EK MADDE 1 (MERKEZİ YERLEŞTİRME PUANI) YATAY GEÇİŞ DEĞERLENDİRME SONUÇLARI'
sub_top.fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
sub_top.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
sub_top.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 24

curr_row = 4

# Dept Banner
ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=len(headers))
dept_cell = ws.cell(row=curr_row, column=1, value='İÇ MİMARLIK VE ÇEVRE TASARIMI BÖLÜMÜ')
dept_cell.fill = dept_banner_fill
dept_cell.font = dept_banner_font
dept_cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
ws.row_dimensions[curr_row].height = 24
curr_row += 1

# Table Headers
for c_idx, h_name in enumerate(headers, start=1):
    cell = ws.cell(row=curr_row, column=c_idx, value=h_name)
    cell.fill = col_header_fill
    cell.font = col_header_font
    cell.border = thick_bottom
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[curr_row].height = 24
curr_row += 1

# Rows
for r_i, r_data in enumerate(rows):
    is_zebra = (r_i % 2 == 1)
    for c_idx, val in enumerate(r_data, start=1):
        cell = ws.cell(row=curr_row, column=c_idx, value=val)
        cell.font = data_font
        cell.border = thin_border
        if is_zebra:
            cell.fill = zebra_fill
        if c_idx in [1, 5, 6, 8, 9, 12, 13]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx in [7, 10]:
            cell.alignment = Alignment(horizontal='right', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[curr_row].height = 24
    curr_row += 1

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 25
ws.column_dimensions['D'].width = 50
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 12
ws.column_dimensions['G'].width = 15
ws.column_dimensions['H'].width = 14
ws.column_dimensions['I'].width = 12
ws.column_dimensions['J'].width = 15
ws.column_dimensions['K'].width = 34
ws.column_dimensions['L'].width = 14
ws.column_dimensions['M'].width = 12

out1 = r'D:\2026 dünya\MTGSF_Ek_Madde_1_Yatay_Gecis_Sonuclari.xlsx'
out2 = r'D:\2026 dünya\kalsın\MTGSF_Ek_Madde_1_Yatay_Gecis_Sonuclari.xlsx'

wb.save(out1)
wb.save(out2)
print('Successfully saved:', out1)
print('Successfully saved:', out2)
