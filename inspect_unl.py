# -*- coding: utf-8 -*-
import openpyxl, json, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('cupmat maç programı.xlsx', data_only=True)
for name in ['U Uluslar L', 'P Durumu']:
    if name in wb.sheetnames:
        ws = wb[name]
        print('=== SHEET: ' + name + ' ===')
        rows = list(ws.iter_rows(values_only=True))
        for r_idx in range(min(30, len(rows))):
            row_data = [str(cell) if cell is not None else '' for cell in rows[r_idx]]
            print('Row ' + str(r_idx) + ': ' + json.dumps(row_data, ensure_ascii=False))
