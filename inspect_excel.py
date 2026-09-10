# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook('cupmat maç programı.xlsx', data_only=True)
print('Sheet Names:')
for i, name in enumerate(wb.sheetnames):
    ws = wb[name]
    print(f'{i+1}. {name} | Rows: {ws.max_row} | Cols: {ws.max_column}')

for name in wb.sheetnames:
    ws = wb[name]
    print(f'\n=== Sheet: {name} ===')
    rows = list(ws.iter_rows(values_only=True))
    for r_idx in range(min(10, len(rows))):
        print(f'Row {r_idx}: {rows[r_idx]}')
